"""
Хранение и чтение файлов библиотеки.

Файлы лежат на диске — `uploads/<user_id>/<uuid>.<ext>`, — а не в базе:
двоичные данные раздули бы PostgreSQL и сломали нынешнюю схему резервных
копий. Имя на диске — идентификатор, исходное имя живёт только в базе:
пользователь вправе назвать файл как угодно, включая пути и управляющие
символы, и это не должно доходить до файловой системы.
"""

from __future__ import annotations

import hashlib
import logging
import os
import re
import uuid
from pathlib import Path
from typing import BinaryIO

log = logging.getLogger('workspace.storage')

ALLOWED = {
    '.pdf': 'application/pdf',
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.txt': 'text/plain',
    '.rtf': 'application/rtf',
}

READ_CHUNK = 1024 * 1024


class UnsupportedFormat(ValueError):
    pass


class DocumentTooLarge(ValueError):
    pass


class ExtractionFailed(RuntimeError):
    pass


def uploads_root() -> Path:
    from config import Config

    root = Path(getattr(Config, 'WORKSPACE_UPLOAD_DIR', 'uploads'))
    root.mkdir(parents=True, exist_ok=True)
    return root


def ext_of(filename: str) -> str:
    return os.path.splitext(filename or '')[1].lower()


def save_upload(file_storage, user_id: int, *, max_mb: int = 50) -> dict:
    """Кладёт файл на диск и возвращает его свойства.

    Размер считается по мере записи, а не по заголовку `Content-Length`:
    заголовок присылает клиент, и верить ему при записи на диск нельзя.
    """
    name = file_storage.filename or ''
    ext = ext_of(name)
    if ext not in ALLOWED:
        raise UnsupportedFormat(
            f'Формат {ext or "без расширения"} не поддерживается. '
            f'Загрузите {", ".join(sorted(ALLOWED))}.'
        )

    limit = max_mb * 1024 * 1024
    folder = uploads_root() / str(user_id)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f'{uuid.uuid4().hex}{ext}'

    digest = hashlib.sha256()
    size = 0
    try:
        with open(path, 'wb') as out:
            while True:
                chunk = file_storage.stream.read(READ_CHUNK)
                if not chunk:
                    break
                size += len(chunk)
                if size > limit:
                    raise DocumentTooLarge(f'Файл больше {max_mb} МБ')
                digest.update(chunk)
                out.write(chunk)
    except Exception:
        # Недописанный файл на диске хуже отсутствующего: он выглядит целым,
        # а прочитать его нельзя.
        path.unlink(missing_ok=True)
        raise

    if size == 0:
        path.unlink(missing_ok=True)
        raise UnsupportedFormat('Файл пустой')

    return {
        'storage_path': str(path.relative_to(uploads_root())),
        'file_size': size,
        'sha256': digest.hexdigest(),
        'mime': ALLOWED[ext],
        'original_filename': os.path.basename(name)[:255],
        'ext': ext,
    }


def absolute(storage_path: str) -> Path:
    """Полный путь к файлу с защитой от выхода за пределы хранилища."""
    root = uploads_root().resolve()
    path = (root / storage_path).resolve()
    if not str(path).startswith(str(root) + os.sep):
        raise ValueError(f'путь вне хранилища: {storage_path}')
    return path


def delete(storage_path: str) -> None:
    try:
        absolute(storage_path).unlink(missing_ok=True)
    except Exception as e:
        log.warning('файл %s не удалён: %s', storage_path, e)


# ------------------------------------------------------------- извлечение


def _pdf_text(path: Path) -> tuple[str, int]:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        pages = len(pdf.pages)
        for page in pdf.pages:
            # Прежняя загрузка в чате обрывала PDF на тридцатой странице и
            # тексте в восемь тысяч знаков, потому что весь текст уходил одним
            # куском в запрос к модели. Здесь документ хранится целиком, а в
            # запрос попадают только нужные фрагменты, — значит, договор на
            # шестьдесят страниц больше не обрывается на середине.
            text = page.extract_text()
            if text:
                parts.append(text)
    return '\n\n'.join(parts), pages


def _docx_text(path: Path) -> tuple[str, int]:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts), 0


def _plain_text(path: Path) -> tuple[str, int]:
    raw = path.read_bytes()
    for encoding in ('utf-8', 'cp1251', 'utf-16'):
        try:
            return raw.decode(encoding), 0
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', errors='replace'), 0


def extract_text(storage_path: str, ext: str | None = None) -> tuple[str, int]:
    """Текст документа и число страниц.

    Пустой результат — не ошибка чтения, а обычное дело: сканированный PDF без
    текстового слоя. Это должно доходить до пользователя как «нужно
    распознавание», а не как сбой.
    """
    path = absolute(storage_path)
    ext = (ext or path.suffix).lower()
    try:
        if ext == '.pdf':
            return _pdf_text(path)
        if ext == '.docx':
            return _docx_text(path)
        return _plain_text(path)
    except Exception as e:
        raise ExtractionFailed(f'не удалось прочитать файл: {e}') from e


def clean_filename_title(filename: str) -> str:
    """Заголовок по умолчанию из имени файла."""
    base = os.path.splitext(os.path.basename(filename or ''))[0]
    base = re.sub(r'[_\-]+', ' ', base).strip()
    base = re.sub(r'\s{2,}', ' ', base)
    return (base[:255] or 'Документ без названия')
