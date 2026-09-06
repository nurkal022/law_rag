"""Рендер дерева в HTML/DOCX/PDF/XLSX на реалистичном договоре поставки."""

import io
import os
import re
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docengine.render import to_docx, to_html, to_pdf, to_xlsx  # noqa: E402
from docengine.render.xlsx import _as_number, _sheet_title  # noqa: E402
from docengine.schema import (  # noqa: E402
    Annex, Clause, Column, DocTree, Meta, Party, Ref, Requisites, Section, Table, renumber,
)

XSS = '<script>alert(1)</script>'


@pytest.fixture()
def tree() -> DocTree:
    supplier = Party(
        role='Поставщик', name='ТОО «Астана Логистик»', kind='legal',
        id_no='123456789012', address='г. Астана, ул. Кунаева, 12, оф. 5',
        phone='+7 701 000 00 00', email='info@astana-log.kz',
        bank='АО «Народный Банк Казахстана»', iban='KZ123456789012345678', bik='HSBKKZKX',
        signatory='Абдрахманов А.А.', signatory_role='Директор', basis='устава',
    )
    buyer = Party(
        role='Покупатель', name='ИП «Сериков»', kind='ip',
        id_no='990101300123', address='г. Алматы, пр. Абая, 150',
        phone='+7 702 111 11 11', email='serikov@mail.kz',
        bank='АО «Kaspi Bank»', iban='KZ876543210987654321', bik='CASPKZKA',
        signatory='Сериков С.С.', signatory_role='Индивидуальный предприниматель',
        basis='свидетельства о регистрации',
    )

    spec = Table(
        id='spec', title='Спецификация товара',
        columns=[
            Column(key='n', title='№', width=0.08),
            Column(key='name', title='Наименование', width=0.42),
            Column(key='qty', title='Кол-во', width=0.15, numeric=True),
            Column(key='sum', title='Сумма, тг', width=0.35, numeric=True),
        ],
        rows=[
            ['1', 'Бумага А4, 500 л.', '100', '150000.00'],
            ['2', 'Картридж лазерный', '10', '350000.00'],
        ],
        total_row=['', 'Итого', '110', '500000.00'],
        note='Цены указаны с учётом НДС.',
    )

    t = DocTree(
        meta=Meta(kind='contract', type_id='supply', title='Договор поставки',
                  subtitle='Поставка канцелярских товаров',
                  legal_basis=[Ref(act='ГК РК', article='406', verified=True)]),
        requisites=Requisites(number='12/2026', city='г. Астана', date='2026-09-07',
                              parties=[supplier, buyer]),
        preamble='ТОО «Астана Логистик», именуемое в дальнейшем «Поставщик», и ИП «Сериков», '
                 'именуемый в дальнейшем «Покупатель», заключили настоящий договор о нижеследующем.',
        sections=[
            Section(key='subject', title='Предмет договора', clauses=[
                Clause(text='Поставщик обязуется передать товар согласно спецификации.',
                       refs=[Ref(act='ГК РК', article='406', verified=True)]),
                Clause(text=f'Ассортимент определяется приложением {XSS}.',
                       subclauses=[Clause(text='Изменение ассортимента оформляется допсоглашением.')]),
            ]),
            Section(key='price', title='Цена и порядок расчётов', clauses=[
                Clause(text='Общая стоимость товара составляет 500 000 (пятьсот тысяч) тенге.'),
                Clause(text='Оплата производится в течение 10 банковских дней.',
                       refs=[Ref(act='ГК РК', article='282', verified=False)]),
            ]),
            Section(key='delivery', title='Порядок поставки', clauses=[
                Clause(text='Поставка осуществляется силами Поставщика.'),
            ]),
            Section(key='liability', title='Ответственность сторон', clauses=[
                Clause(text='За просрочку оплаты начисляется пеня 0,1% за каждый день.',
                       refs=[Ref(act='ГК РК', article='353', verified=True)]),
            ]),
            Section(key='final', title='Заключительные положения', clauses=[
                Clause(text='Договор вступает в силу с момента подписания.'),
            ]),
        ],
        tables=[spec],
        annexes=[Annex(title='Спецификация', kind='table', table=spec)],
    )
    return renumber(t)


# ------------------------------------------------------------------ HTML

def test_html_fragment_has_content(tree):
    html = to_html(tree)
    assert 'Договор поставки' in html
    assert '1. ПРЕДМЕТ ДОГОВОРА' in html
    assert '>1.1.<' in html or '1.1.' in html
    assert '2.2.' in html
    assert 'ТОО «Астана Логистик»' in html
    assert 'ИП «Сериков»' in html
    assert 'KZ876543210987654321' in html
    assert 'Спецификация товара' in html
    # фрагмент — без обвязки страницы, чтобы вставлялся в React-предпросмотр
    assert '<!DOCTYPE html>' not in html
    assert html.lstrip().startswith('<div class="doc-sheet"')


def test_html_escapes_user_text(tree):
    html = to_html(tree)
    assert '<script>' not in html
    assert '&lt;script&gt;' in html


def test_html_refs_are_marked(tree):
    html = to_html(tree)
    assert 'data-act="ГК РК"' in html
    assert 'data-article="406"' in html
    assert 'ГК РК ст. 406' in html
    # невыверенная ссылка помечена отдельно и не выдаётся за проверенную
    assert 'ref ref--unverified' in html
    assert re.search(r'class="ref ref--unverified"[^>]*data-article="282"', html)


def test_html_standalone_is_full_page(tree):
    html = to_html(tree, standalone=True)
    assert html.startswith('<!DOCTYPE html>')
    assert '@page' in html and '20mm' in html
    assert 'Times New Roman' in html
    assert '</html>' in html.strip()[-20:]


# ------------------------------------------------------------------ DOCX

def test_docx_bytes_and_filename(tree):
    data, name = to_docx(tree)
    assert data[:2] == b'PK'
    assert len(data) > 5000
    assert name == 'Dogovor_postavki_2026-09-07.docx'
    assert name.isascii()


def test_docx_structure(tree):
    import docx

    data, _ = to_docx(tree)
    doc = docx.Document(io.BytesIO(data))

    texts = [p.text for p in doc.paragraphs]
    assert any('ДОГОВОР ПОСТАВКИ' in t for t in texts)
    assert any(t.startswith('1.1.') for t in texts)
    assert any('1. ПРЕДМЕТ ДОГОВОРА' in t for t in texts)
    # разделы — стилями, а не ручным жирным: иначе не работает навигация Word
    assert any(p.style.name.startswith('Heading') and 'ПРЕДМЕТ' in p.text for p in doc.paragraphs)
    # ссылка на норму дошла до текста пункта
    assert any('ГК РК ст. 406' in t for t in texts)

    # спецификация — настоящая таблица со всеми строками и жирной шапкой
    spec = [t for t in doc.tables if t.cell(0, 0).text == '№']
    assert spec, [t.cell(0, 0).text for t in doc.tables]
    spec = spec[0]
    assert spec.cell(0, 1).text == 'Наименование'
    assert spec.cell(0, 0).paragraphs[0].runs[0].bold is True  # шапка жирная
    assert 'Картридж лазерный' in spec.cell(2, 1).text
    assert 'Итого' in spec.rows[-1].cells[1].text

    # блок подписей — таблица в две колонки с реквизитами обеих сторон
    sign = [t for t in doc.tables if 'ПОСТАВЩИК' in t.cell(0, 0).text]
    assert sign, 'нет таблицы подписей'
    sign = sign[0]
    assert len(sign.columns) == 2
    assert 'ПОКУПАТЕЛЬ' in sign.cell(0, 1).text
    assert 'KZ123456789012345678' in sign.cell(0, 0).text
    assert 'Абдрахманов А.А.' in sign.cell(0, 0).text
    assert 'Сериков С.С.' in sign.cell(0, 1).text

    # приложение
    assert any('Приложение № 1' in t for t in texts)


def test_docx_page_setup(tree):
    import docx
    data, _ = to_docx(tree)
    doc = docx.Document(io.BytesIO(data))
    sec = doc.sections[0]
    assert round(sec.left_margin.mm) == 20 and round(sec.top_margin.mm) == 20
    assert round(sec.right_margin.mm) == 20 and round(sec.bottom_margin.mm) == 20
    assert doc.styles['Normal'].font.name == 'Times New Roman'
    assert doc.styles['Normal'].font.size.pt == 12
    footer_text = ' '.join(p.text for p in sec.footer.paragraphs)
    assert 'Договор поставки' in footer_text


# ------------------------------------------------------------------ XLSX

def test_xlsx_sheets_and_header(tree):
    import openpyxl

    data, name = to_xlsx(tree)
    assert name == 'Dogovor_postavki_2026-09-07.xlsx'
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert 'Спецификация товара' in wb.sheetnames
    ws = wb['Спецификация товара']
    assert [c.value for c in ws[1]] == ['№', 'Наименование', 'Кол-во', 'Сумма, тг']
    assert all(c.font.bold for c in ws[1])
    assert ws['A1'].fill.fgColor.rgb.endswith('DCE6F1')
    # числовые колонки — числами, чтобы по выгрузке считались формулы
    assert ws['D2'].value == 150000.0
    assert ws['D2'].number_format == '#,##0.00'
    # строка «Итого» жирная
    total = ws[4]  # 1 шапка + 2 строки + итог
    assert any('Итого' in str(c.value or '') for c in total)
    assert all(c.font.bold for c in total)
    assert ws['C4'].value == 110
    assert ws.column_dimensions['B'].width > 10


def test_xlsx_single_table(tree):
    import openpyxl

    data, _ = to_xlsx(tree, table_id='spec')
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert len(wb.sheetnames) == 1

    with pytest.raises(ValueError):
        to_xlsx(tree, table_id='nope')


def test_xlsx_sheet_title_sanitized():
    used = set()
    long = 'Спецификация [товара]: поставка/отгрузка по договору № 12'
    name = _sheet_title(long, used)
    assert len(name) <= 31
    assert not set(name) & set('\\/*?:[]')
    assert _sheet_title(long, used) != name  # дубликаты разводятся


@pytest.mark.parametrize('raw,expected', [
    ('150000.00', 150000.0), ('1 200,50', 1200.5), ('110', 110),
    ('по договорённости', None), ('', None),
])
def test_xlsx_number_parsing(raw, expected):
    assert _as_number(raw) == expected


# ------------------------------------------------------------------ PDF

def test_pdf(tree):
    try:
        import weasyprint  # noqa: F401
    except (ImportError, OSError) as exc:
        pytest.skip(f'WeasyPrint недоступен в этом окружении ({exc}); PDF не проверяем')
    data, name = to_pdf(tree)
    assert data[:4] == b'%PDF'
    assert name.endswith('.pdf') and name.isascii()
