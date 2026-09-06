"""Документ в DOCX.

Ключевое отличие от старого экспорта: это настоящий Word-документ — заголовки
разделов сделаны стилями (значит, работают навигация и оглавление в Word),
таблицы — таблицами, а не текстом с пробелами. Юрист дорабатывает выгрузку
руками, и от того, останется ли структура структурой, зависит вся дальнейшая
работа с файлом.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from ..schema import Annex, DocTree, Party, Section, Table
from .html import fmt_date
from .naming import filename

FONT = 'Times New Roman'
SIZE = Pt(12)


def _force_font(run_font, name: str = FONT) -> None:
    """python-docx проставляет только w:ascii/w:hAnsi.

    Кириллица в Word берётся из hAnsi, но при копировании текста из других
    редакторов может уехать в eastAsia/cs — поэтому прибиваем все четыре слота.
    """
    run_font.name = name
    rpr = run_font.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), name)


def _setup_styles(doc: Document) -> None:
    normal = doc.styles['Normal']
    _force_font(normal.font)
    normal.font.size = SIZE
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    # Штатные Heading в Word — синий Calibri Light; для договора это чужеродно,
    # но сам стиль сохраняем: на нём держатся навигация и автооглавление.
    for name, size in (('Heading 1', 13), ('Heading 2', 12), ('Title', 14)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _force_font(st.font)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Mm(20)
        sec.top_margin = sec.bottom_margin = Mm(20)


def _page_field(paragraph) -> None:
    """Номер страницы — полем PAGE, а не текстом: иначе он не пересчитается."""
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, end):
        run._r.append(el)


def _footer(doc: Document, title: str) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{title} — стр. ')
    run.font.size = Pt(9)
    _force_font(run.font)
    _page_field(p)
    for r in p.runs:
        r.font.size = Pt(9)


def _para(doc: Document, text: str, *, align=None, indent=None, bold=False,
          size: int | None = None, space_after=Pt(4)):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.first_line_indent = indent
    p.paragraph_format.space_after = space_after
    run = p.add_run(text)
    run.bold = bold
    _force_font(run.font)
    if size:
        run.font.size = Pt(size)
    return p


def _clause_text(clause) -> str:
    """Ссылки на нормы уходят в текст пункта: в Word нет места для разметки,
    а юристу важно видеть основание прямо в абзаце."""
    if not clause.refs:
        return clause.text
    return f"{clause.text} ({'; '.join(r.label() for r in clause.refs)})"


def _sections(doc: Document, sections: list[Section], heading_level: int = 1) -> None:
    for s in sections:
        h = doc.add_heading(f'{s.no}. {s.title.upper()}', level=heading_level)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in h.runs:
            _force_font(run.font)
        if s.pending:
            _para(doc, '[Раздел не сгенерирован]', indent=Cm(1.25))
        for c in s.clauses:
            p = _para(doc, f'{c.no}. {_clause_text(c)}', indent=Cm(1.25))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for sub in c.subclauses:
                sp = _para(doc, f'{sub.no}. {_clause_text(sub)}', indent=Cm(0.75))
                sp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                sp.paragraph_format.left_indent = Cm(1.25)


def _table(doc: Document, t: Table) -> None:
    _para(doc, t.title, bold=True, space_after=Pt(2))
    rows = len(t.rows) + (1 if t.total_row else 0) + 1
    table = doc.add_table(rows=rows, cols=len(t.columns))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    # шапка повторяется на каждой странице длинной таблицы
    tr_pr = hdr._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:tblHeader'))
    for i, col in enumerate(t.columns):
        cell = hdr.cells[i]
        run = cell.paragraphs[0].add_run(col.title)
        run.bold = True
        _force_font(run.font)
        run.font.size = Pt(11)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    body_rows = list(t.rows) + ([t.total_row] if t.total_row else [])
    for ri, row in enumerate(body_rows, start=1):
        is_total = t.total_row is not None and ri == len(body_rows)
        for ci, col in enumerate(t.columns):
            cell = table.rows[ri].cells[ci]
            value = row[ci] if ci < len(row) else ''
            run = cell.paragraphs[0].add_run(value)
            run.bold = is_total
            _force_font(run.font)
            run.font.size = Pt(11)
            if col.numeric:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    if t.note:
        _para(doc, t.note, size=10)


def _party_lines(p: Party) -> list[str]:
    lines = [p.role.upper(), p.name]
    if p.id_no:
        lines.append(f"{'БИН' if p.kind == 'legal' else 'ИИН'}: {p.id_no}")
    if p.address:
        lines.append(f'Адрес: {p.address}')
    if p.phone:
        lines.append(f'Тел.: {p.phone}')
    if p.email:
        lines.append(f'E-mail: {p.email}')
    if p.bank:
        lines.append(f'Банк: {p.bank}')
    if p.iban:
        lines.append(f'ИИК: {p.iban}')
    if p.bik:
        lines.append(f'БИК: {p.bik}')
    lines.append('')
    lines.append(f'{p.signatory_role} _________________ {p.signatory}'.strip())
    if p.basis:
        lines.append(f'Действует на основании: {p.basis}')
    return [line for line in lines if line is not None]


def _signatures(doc: Document, parties: list[Party]) -> None:
    h = doc.add_heading('РЕКВИЗИТЫ И ПОДПИСИ СТОРОН', level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in h.runs:
        _force_font(run.font)

    # Две колонки — таблица без стиля: Word не рисует границы у таблицы
    # без табличного стиля, а колонки держат реквизиты сторон друг напротив друга.
    cols = 2
    rows = (len(parties) + cols - 1) // cols
    table = doc.add_table(rows=rows or 1, cols=cols)
    table.autofit = True
    for idx, party in enumerate(parties):
        cell = table.rows[idx // cols].cells[idx % cols]
        first = True
        for line in _party_lines(party):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            _force_font(run.font)
            run.font.size = Pt(10)
            if line == party.role.upper() or line == party.name:
                run.bold = True


def _annexes(doc: Document, annexes: list[Annex]) -> None:
    for a in annexes:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _para(doc, f'Приложение № {a.no}', align=WD_ALIGN_PARAGRAPH.RIGHT)
        t = doc.add_heading(a.title.upper(), level=1)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in t.runs:
            _force_font(run.font)
        if a.kind == 'table' and a.table is not None:
            _table(doc, a.table)
        else:
            _sections(doc, a.sections)


def to_docx(tree: DocTree) -> tuple[bytes, str]:
    doc = Document()
    _setup_styles(doc)
    _footer(doc, tree.meta.title)

    title = doc.add_heading(tree.meta.title.upper(), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        _force_font(run.font)
        run.font.size = Pt(14)
        run.font.bold = True
    if tree.requisites.number:
        _para(doc, f'№ {tree.requisites.number}', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    if tree.meta.subtitle:
        _para(doc, tree.meta.subtitle, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Город слева, дата справа — одной строкой через таблицу без границ:
    # табуляции разъезжаются при смене шрифта.
    place = doc.add_table(rows=1, cols=2)
    left, right = place.rows[0].cells
    lr = left.paragraphs[0].add_run(tree.requisites.city)
    _force_font(lr.font)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = right.paragraphs[0].add_run(fmt_date(tree.requisites.date))
    _force_font(rr.font)
    doc.add_paragraph()

    if tree.preamble:
        p = _para(doc, tree.preamble, indent=Cm(1.25), space_after=Pt(8))
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _sections(doc, tree.sections)

    for t in tree.tables:
        _table(doc, t)

    if tree.requisites.parties:
        _signatures(doc, tree.requisites.parties)

    _annexes(doc, tree.annexes)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), filename(tree, 'docx')
